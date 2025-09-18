#!/usr/bin/env python3
"""
Тест импорта Excel для отладки проблемы
"""

import requests
import json
import os
from datetime import datetime

BASE_URL = "http://127.0.0.1:8000/api/v1"

def get_auth_token():
    """Получить токен аутентификации"""
    response = requests.post(f"{BASE_URL}/auth/token/", json={
        "username": "admin",
        "password": "password123"
    })
    if response.status_code == 200:
        return response.json()["access"]
    else:
        print(f"Ошибка аутентификации: {response.status_code}")
        return None

def create_test_excel():
    """Создать тестовый Excel файл"""
    from openpyxl import Workbook
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Закупки"
    
    # Заголовки
    headers = [
        "Дата", "Объект", "Ответственный", "Поставщик", "Комментарий",
        "Материал", "Количество", "Единица", "Цена за единицу", "Сумма"
    ]
    
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Тестовые данные
    test_data = [
        ["2025-09-17", "ЖК 'Зеленый парк' - Блок А", "Админ Админов", "ООО Стройматериалы", "Тестовая закупка",
         "Цемент М400", "100", "кг", "50.00", "5000.00"],
        ["2025-09-17", "ЖК 'Зеленый парк' - Блок А", "Админ Админов", "ООО Стройматериалы", "Тестовая закупка",
         "Песок речной", "200", "кг", "25.00", "5000.00"]
    ]
    
    for row, data in enumerate(test_data, 2):
        for col, value in enumerate(data, 1):
            ws.cell(row=row, column=col, value=value)
    
    filename = "test_import.xlsx"
    wb.save(filename)
    return filename

def test_import_prepare():
    """Тестирование подготовки импорта"""
    token = get_auth_token()
    if not token:
        return False
    
    headers = {"Authorization": f"Bearer {token}"}
    
    print("🧪 Тестирование импорта Excel...")
    
    # Создаем тестовый файл
    filename = create_test_excel()
    print(f"📄 Создан тестовый файл: {filename}")
    
    # Тестируем подготовку импорта
    with open(filename, 'rb') as f:
        files = {'file': (filename, f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        
        print("📤 Отправляем файл на подготовку...")
        response = requests.post(
            f"{BASE_URL}/purchases/import/prepare/",
            headers=headers,
            files=files
        )
    
    print(f"📊 Статус ответа: {response.status_code}")
    print(f"📄 Ответ: {response.text[:500]}")
    
    # Очищаем тестовый файл
    if os.path.exists(filename):
        os.remove(filename)
    
    if response.status_code == 200:
        print("✅ Подготовка импорта работает!")
        return True
    else:
        print("❌ Ошибка при подготовке импорта")
        return False

def test_import_endpoints():
    """Тестирование всех endpoints импорта"""
    token = get_auth_token()
    if not token:
        return False
    
    headers = {"Authorization": f"Bearer {token}"}
    
    print("\n🔍 Проверяем доступность endpoints импорта...")
    
    endpoints = [
        ("/purchases/import/prepare/", "Подготовка"),
        ("/purchases/import/dry_run/", "Сухой прогон"),
        ("/purchases/import/commit/", "Коммит")
    ]
    
    for endpoint, name in endpoints:
        # Простой GET запрос для проверки доступности
        response = requests.get(f"{BASE_URL}{endpoint}", headers=headers)
        print(f"  {name}: {response.status_code} ({'✅' if response.status_code in [200, 405] else '❌'})")
    
    return True

def main():
    """Основная функция тестирования"""
    print("🚀 Тестирование импорта Excel")
    print("=" * 50)
    
    # Тест endpoints
    endpoints_ok = test_import_endpoints()
    
    # Тест подготовки импорта
    prepare_ok = test_import_prepare()
    
    print("\n" + "=" * 50)
    print("📋 Результаты:")
    print(f"  Endpoints: {'✅' if endpoints_ok else '❌'}")
    print(f"  Подготовка импорта: {'✅' if prepare_ok else '❌'}")
    
    return endpoints_ok and prepare_ok

if __name__ == "__main__":
    main()
